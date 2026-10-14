import unittest
from io import BytesIO
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from converter import convert
from contrast import contrast_ratio, minimum_ratio

class ContrastTests(unittest.TestCase):
    def test_reference_ratios(self):
        self.assertEqual(contrast_ratio('000000','FFFFFF'),21)
        self.assertAlmostEqual(contrast_ratio('777777','FFFFFF'),4.478,places=3)
        self.assertLess(contrast_ratio('CDCDCD','FFFFFF'),4.5)
        self.assertEqual(minimum_ratio(18),3)
        self.assertEqual(minimum_ratio(14,True),3)
        self.assertEqual(minimum_ratio(14,False),4.5)

    def test_light_text_and_dark_cell_are_fixed(self):
        doc=Document();p=doc.add_paragraph('Light text');p.runs[0].font.color.rgb=RGBColor(205,205,205)
        table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='Dark cell'
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'000000');table.cell(0,0)._tc.get_or_add_tcPr().append(shade)
        doc.sections[0].header.paragraphs[0].add_run('Light header').font.color.rgb=RGBColor(220,220,220)
        buf=BytesIO();doc.save(buf)
        output,report=convert(buf.getvalue())
        fixed=Document(BytesIO(output))
        for root in [fixed.element,fixed.sections[0].header._element]:
            for run in root.xpath('.//w:r[w:t]'):
                self.assertEqual(run.xpath('./w:rPr/w:color')[0].get(qn('w:val')),'000000')
                self.assertEqual(run.xpath('./w:rPr/w:shd')[0].get(qn('w:fill')),'FFFFFF')
        self.assertEqual(fixed.tables[0].cell(0,0)._tc.xpath('./w:tcPr/w:shd')[0].get(qn('w:fill')),'FFFFFF')
        self.assertIn('21:1',report)
        self.assertIn('below 4.5:1',report)
        self.assertIn('unsupported parts',report)
        original,_=convert(buf.getvalue(),high_contrast=False)
        self.assertEqual(str(Document(BytesIO(original)).paragraphs[0].runs[0].font.color.rgb),'CDCDCD')
