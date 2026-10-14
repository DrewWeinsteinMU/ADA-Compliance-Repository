import unittest
from io import BytesIO
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from documents import process_document


def pdf_sample(text=True):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 72 720 Td (Employee document example) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
    buffer = BytesIO(); writer.write(buffer)
    return buffer.getvalue()

class DocumentTests(unittest.TestCase):
    def test_word_html_order_and_escaping(self):
        doc=Document(); doc.add_paragraph('Employee Guide'); doc.add_heading('Instructions',1)
        doc.add_paragraph('<script>alert("x")</script>')
        table=doc.add_table(rows=2,cols=1); table.cell(0,0).text='Column'; table.cell(1,0).text='Record'
        doc.add_paragraph('After table')
        data=BytesIO(); doc.save(data)
        output=process_document(data.getvalue(),'guide.docx',table_headers=True)
        html=output['html'].decode()
        self.assertIn('<h2>Instructions</h2>',html)
        self.assertNotIn('<script>',html)
        self.assertIn('&lt;script&gt;',html)
        self.assertIn('<th scope="col">Column</th>',html)
        self.assertLess(html.index('Record'),html.index('After table'))
        self.assertIn('lang="en-US"',html)

    def test_pdf_extraction_and_scan(self):
        output=process_document(pdf_sample(),'employee.pdf')
        self.assertIn(b'Employee document example',output['html'])
        self.assertIn(b'source PDF is not repaired in place',output['report'])
        self.assertIn(b'missing an identifiable marked structure tree',output['report'])
        doc=Document(BytesIO(output['document']))
        self.assertIn('Employee document example',' '.join(p.text for p in doc.paragraphs))

    def test_pdf_without_text_requires_ocr(self):
        with self.assertRaisesRegex(ValueError,'OCR'):
            process_document(pdf_sample(False),'scan.pdf')

    def test_encrypted_pdf(self):
        writer=PdfWriter(); writer.add_blank_page(width=100,height=100); writer.encrypt('secret')
        data=BytesIO(); writer.write(data)
        with self.assertRaisesRegex(ValueError,'Password-protected'):
            process_document(data.getvalue(),'locked.pdf')

    def test_invalid_extension_and_content(self):
        with self.assertRaises(ValueError): process_document(b'%PDF-broken','bad.pdf')
        with self.assertRaises(ValueError): process_document(b'any','file.exe')
