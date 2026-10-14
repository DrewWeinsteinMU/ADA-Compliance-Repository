import unittest
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from converter import convert

class ConverterTests(unittest.TestCase):
    def sample(self):
        doc=Document()
        doc.add_paragraph('Sample Document')
        doc.add_paragraph('Second Title')
        doc.add_paragraph('Body content stays unchanged.')
        table=doc.add_table(rows=2, cols=2)
        table.cell(0,0).text='Name'; table.cell(0,1).text='Value'
        table.cell(1,0).text='Example'; table.cell(1,1).text='42'
        data=BytesIO(); doc.save(data)
        return data.getvalue()

    def test_preservation_and_headings(self):
        out,report=convert(self.sample())
        doc=Document(BytesIO(out))
        self.assertEqual([p.text for p in doc.paragraphs],['Sample Document','Second Title','Body content stays unchanged.'])
        self.assertEqual(doc.paragraphs[1].style.name,'Heading 1')
        self.assertEqual(doc.tables[0].cell(1,1).text,'42')
        self.assertIn('Table 1:',report)
        self.assertIn('does not certify',report)
        with ZipFile(BytesIO(out)) as z: self.assertIsNone(z.testzip())

    def test_options(self):
        out,_=convert(self.sample(),detect_headings=False,table_headers=True)
        doc=Document(BytesIO(out))
        self.assertEqual(doc.paragraphs[1].style.name,'Normal')
        self.assertTrue(doc.tables[0].rows[0]._tr.xpath('./w:trPr/w:tblHeader'))

    def test_reject_invalid(self):
        with self.assertRaises(ValueError): convert(b'not a docx')

if __name__=='__main__': unittest.main()
