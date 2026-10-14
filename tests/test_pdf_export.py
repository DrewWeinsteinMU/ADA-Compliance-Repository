import shutil
import unittest
from io import BytesIO
from docx import Document
from pypdf import PdfReader
from pdf_export import export_tagged_pdf

class TaggedPDFTests(unittest.TestCase):
    @unittest.skipUnless(shutil.which('soffice') or shutil.which('libreoffice'), 'LibreOffice required')
    def test_structure_heading_table_and_text(self):
        doc=Document(); doc.core_properties.title='Employee Guide'
        doc.add_heading('Instructions',1); doc.add_paragraph('Read this accessible employee document.')
        table=doc.add_table(rows=1,cols=2); table.cell(0,0).text='Employee'; table.cell(0,1).text='Example'
        source=BytesIO();doc.save(source)
        data=export_tagged_pdf(source.getvalue())
        pdf=PdfReader(BytesIO(data)); root=pdf.trailer['/Root']
        self.assertTrue(root['/MarkInfo']['/Marked'])
        self.assertTrue(root.get('/Lang'))
        self.assertIn('Read this accessible employee document',pdf.pages[0].extract_text())
        roles=[]
        def walk(node):
            if hasattr(node,'get_object'):node=node.get_object()
            if isinstance(node,list):
                for child in node:walk(child)
            elif isinstance(node,dict):
                if '/S' in node:roles.append(str(node['/S']))
                if '/K' in node:walk(node['/K'])
        walk(root['/StructTreeRoot'])
        self.assertIn('/H1',roles)
        self.assertIn('/Table',roles)
        self.assertIn('/TD',roles)
