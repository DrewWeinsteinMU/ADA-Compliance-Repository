import json
import threading
import unittest
from io import BytesIO
from http.server import ThreadingHTTPServer
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from docx import Document
from app import Handler, RESULTS, LOCK
from test_documents import pdf_sample

class AppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True); cls.thread.start()
        cls.root='http://127.0.0.1:'+str(cls.server.server_port)

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown(); cls.server.server_close(); cls.thread.join()
        with LOCK: RESULTS.clear()

    def upload(self,data,name):
        boundary='test-upload-boundary'
        body=(f'--{boundary}\r\nContent-Disposition: form-data; name="document"; filename="{name}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode()+data+f'\r\n--{boundary}--\r\n'.encode())
        return urlopen(Request(self.root+'/convert',body,{'Content-Type':'multipart/form-data; boundary='+boundary}))

    def test_word_and_pdf_downloads(self):
        doc=Document(); doc.add_paragraph('Employee content'); buf=BytesIO(); doc.save(buf)
        for content,name in [(buf.getvalue(),'employee.docx'),(pdf_sample(),'employee.pdf')]:
            with self.subTest(name=name):
                result=json.load(self.upload(content,name))
                self.assertTrue(result['review_supported'])
                token=result['token']
                review=json.load(urlopen(self.root+'/review/'+token))
                self.assertIn('headings',review)
                self.assertIn('images',review)
                word=urlopen(self.root+'/download/'+token+'/document')
                self.assertIn('Accessible_Document.docx',word.headers['Content-Disposition'])
                self.assertTrue(Document(BytesIO(word.read())).paragraphs)
                html=urlopen(self.root+'/download/'+token+'/html')
                self.assertIn('Document_Text.html',html.headers['Content-Disposition'])
                self.assertIn(b'<main>',html.read())
                self.assertIn(b'HUMAN REVIEW REQUIRED',urlopen(self.root+'/download/'+token+'/report').read())

    def test_invalid_upload_is_actionable(self):
        with self.assertRaises(HTTPError) as context: self.upload(b'not word','bad.docx')
        self.assertEqual(context.exception.code,400)
        self.assertIn('error',json.load(context.exception))

    def test_expired_download(self):
        with LOCK: RESULTS['expired']=(0,{})
        with self.assertRaises(HTTPError) as context: urlopen(self.root+'/download/expired/document')
        self.assertEqual(context.exception.code,404)

class VersionTests(unittest.TestCase):
    def test_review_capability_and_real_expiration(self):
        server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        root='http://127.0.0.1:'+str(server.server_port)
        try:
            health=json.load(urlopen(root+'/health'))
            self.assertTrue(health['review_supported'])
            self.assertTrue(health['version'])
            with self.assertRaises(HTTPError) as context:urlopen(root+'/review/missing-token')
            self.assertEqual(json.load(context.exception)['error'],'Review expired. Upload the document again.')
        finally:
            server.shutdown();server.server_close();thread.join()
