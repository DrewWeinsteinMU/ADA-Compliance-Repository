import unittest
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from review import manifest,apply_review
from converter import convert
from pathlib import Path

class ReviewTests(unittest.TestCase):
    def sample(self):return Path('tests/fixtures/Accessibility_Stress_Test.docx').read_bytes()
    def test_full_review_preserves_content_and_applies_edits(self):
        data,_=convert(self.sample())
        before=Document(BytesIO(data));m=manifest(data)
        self.assertEqual(len(m['images']),2)
        self.assertTrue(m['images'][0]['preview'].startswith('data:image/png;base64,'))
        self.assertEqual(len(m['tables']),4)
        edits={
          'headings':[{'id':h['id'],'level':h['suggested']} for h in m['headings']],
          'images':[{'id':i['id'],'alt':'Training sessions: 2, 4, and 6.'} for i in m['images']],
          'links':[{'id':l['id'],'label':'Example employee resource'} for l in m['links']],
          'tables':[{'id':t['id'],'action':'simplify' if t['complex'] else 'header','confirmed':True} for t in m['tables']]}
        result,changes=apply_review(data,edits);after=Document(BytesIO(result))
        self.assertEqual(len(after.tables),2)
        self.assertEqual(len(after.inline_shapes),len(before.inline_shapes))
        self.assertTrue(all(i['alt'] for i in manifest(result)['images']))
        self.assertTrue(all(t['header'] for t in manifest(result)['tables']))
        self.assertTrue(all(l['label']=='Example employee resource' for l in manifest(result)['links']))
        self.assertIn('Training notes for all departments',' '.join(p.text for p in after.paragraphs))
        self.assertTrue(changes)
        self.assertTrue(all(h['level']==h['suggested'] for h in manifest(result)['headings']))
        sizes=[int(e.get(qn('w:val'))) for e in after.element.xpath('.//w:r[w:t]/w:rPr/w:sz')]
        self.assertTrue(all(size>=24 for size in sizes))

    def test_simplify_requires_confirmation_and_indices_are_validated(self):
        data,_=convert(self.sample())
        with self.assertRaisesRegex(ValueError,'Confirm'):apply_review(data,{'tables':[{'id':1,'action':'simplify'}]})
        with self.assertRaises(ValueError):apply_review(data,{'images':[{'id':999,'alt':'text'}]})
        with self.assertRaises(ValueError):apply_review(data,{'headings':[{'id':0,'level':9}]})
        with self.assertRaises(ValueError):apply_review(data,{'links':[{'id':0,'label':''}]})
