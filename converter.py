"""Conservative Word accessibility remediation; no compliance certification."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MAX_UPLOAD = 15 * 1024 * 1024


def convert(data, detect_headings=True, table_headers=False, high_contrast=True, enlarge_text=True):
    if len(data) > MAX_UPLOAD:
        raise ValueError('Please upload a Word document smaller than 15 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError('This document expands beyond the supported size limit.')
            if any(i.flag_bits & 1 for i in archive.infolist()):
                raise ValueError('Password-protected documents are not supported.')
            if 'word/document.xml' not in archive.namelist():
                raise ValueError('Please upload a valid .docx document.')
        doc = Document(BytesIO(data))
    except (BadZipFile, KeyError, TypeError) as exc:
        raise ValueError('This file could not be read. Save it as .docx in Word and try again.') from exc
    changes, review = [], []
    small_text = sum(1 for p in doc.paragraphs for r in p.runs if r.text.strip() and r.font.size and r.font.size.pt < 11)
    if small_text:
        review.append(f'{small_text} directly formatted text run(s) are smaller than 11 pt. Review and enlarge where appropriate.')
    levels = [int(p.style.name.split()[-1]) for p in doc.paragraphs if p.style.name.startswith('Heading ') and p.style.name.split()[-1].isdigit()]
    if not levels:
        review.append('The source has no semantic section headings. Review the suggested headings or apply heading styles in Word.')
    previous = 0
    for level in levels:
        if level > previous + 1:
            review.append('The source skips heading levels. Review the hierarchy in the Navigation Pane.')
            break
        previous = level
    for name, size in [('Normal', 12), ('Title', 24), ('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12)]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.15
        if name.startswith('Heading'):
            style.paragraph_format.keep_with_next = True
    changes.append('Updated standard text and heading styles for readable type, black text, and consistent spacing.')
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    if defaults is None:
        defaults = OxmlElement('w:docDefaults'); doc.styles.element.insert(0, defaults)
    rdefaults = defaults.find(qn('w:rPrDefault'))
    if rdefaults is None:
        rdefaults = OxmlElement('w:rPrDefault'); defaults.append(rdefaults)
    rpr = rdefaults.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr'); rdefaults.append(rpr)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'en-US'); rpr.append(lang)
        changes.append('Set the default proofing language to English (US); existing explicit language tags were retained.')
        review.append('Confirm the default language is appropriate and mark passages in other languages in Word.')
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    inferred = []
    if detect_headings:
        for index, p in enumerate(paragraphs):
            text = p.text.strip()
            if p.style.name != 'Normal' or len(text) > 100 or len(text.split()) > 12:
                continue
            if p._p.xpath('.//w:drawing | .//w:pict | .//w:fldChar | .//w:hyperlink | .//w:numPr'):
                continue
            runs = [r for r in p.runs if r.text.strip()]
            is_label = text.lower().endswith('title') or text.lower().startswith('title of ') or text.isupper()
            is_bold = bool(runs) and all(r.bold for r in runs)
            if index == 0 and not text.endswith(('.', '?', '!')):
                p.style = 'Title'; inferred.append(text)
            elif (is_label or is_bold) and not text.endswith(('.', '?', '!')):
                p.style = 'Heading 1'; inferred.append(text)
    if inferred:
        changes.append('Assigned title or Heading 1 styles to likely headings: ' + '; '.join(inferred))
        review.append('Check inferred headings in the Navigation Pane. Adjust levels to match the actual document hierarchy.')
    if not doc.core_properties.title:
        doc.core_properties.title = paragraphs[0].text[:200] if paragraphs else 'Accessible document'
        changes.append('Added document title metadata.')
    doc.core_properties.language = doc.core_properties.language or 'en-US'
    roots = [doc.element]
    for section in doc.sections:
        roots.extend([section.header._element, section.footer._element])
    missing_alt = sum(1 for root in roots for e in root.xpath('.//*[local-name()="docPr"]') if not (e.get('descr') or '').strip())
    if missing_alt:
        review.append(f'{missing_alt} image/object instance(s) lack description text. Add meaningful alt text or mark decorative objects in Word.')
    if any(root.xpath('.//w:pict') for root in roots):
        review.append('Legacy image objects are present. Review their alternative text and reading order in Word.')
    for link in doc.element.xpath('.//w:hyperlink'):
        label = ''.join(link.xpath('.//w:t/text()')).strip().lower()
        if label in ('click here', 'here', 'read more', 'more', 'link') or not label:
            review.append(f'WCAG 2.4.4: hyperlink label "{label or "(empty)"}" needs contextual review and descriptive wording.')
    tables = doc.element.xpath('.//w:tbl')
    for index, table in enumerate(tables, 1):
        rows = table.findall(qn('w:tr'))
        if table_headers and rows:
            trpr = rows[0].find(qn('w:trPr'))
            if trpr is None:
                trpr = OxmlElement('w:trPr'); rows[0].insert(0, trpr)
            header = trpr.find(qn('w:tblHeader'))
            if header is None:
                header = OxmlElement('w:tblHeader'); trpr.append(header)
            header.set(qn('w:val'), 'true')
        elif rows and not rows[0].xpath('./w:trPr/w:tblHeader[not(@w:val="0") and not(@w:val="false")]'):
            review.append(f'Table {index}: confirm whether the first row contains column headings, then designate a header row in Word.')
        if table.xpath('.//w:gridSpan | .//w:vMerge | .//w:tbl/w:tr/w:tc/w:tbl'):
            review.append(f'Table {index}: merged cells or nested tables require manual review for a clear reading order.')
    if table_headers and tables:
        changes.append(f'Marked the first row of {len(tables)} table(s) as repeating headers, as requested.')
    if doc.element.xpath('.//w:txbxContent'):
        review.append('Text boxes are present. Check reading order and move essential text into the main document if needed.')
    if doc.element.xpath('.//w:ins | .//w:del'):
        review.append('Tracked changes are present. Review and resolve them before distributing the document.')
    if enlarge_text:
        from review import minimum_text
        enlarged = minimum_text(doc)
        changes.append(f'Raised {enlarged} directly formatted text runs below 12 points to 12 points. This is a readability preference, not a WCAG minimum font-size rule.')
    if high_contrast:
        from contrast import normalize_contrast
        contrast_changes, contrast_review = normalize_contrast(doc)
        changes.extend(contrast_changes); review.extend(contrast_review)
    else:
        review.append('High-contrast correction was disabled. Original direct colors and backgrounds require review against WCAG 2.2 SC 1.4.3.')
    review.extend(['Review descriptive link text and reading order. Direct font sizes and other formatting may still require review.',
                   'Run Review → Check Accessibility in Microsoft Word, and review the document with a screen reader before distribution.'])
    output = BytesIO(); doc.save(output)
    report = 'WORD ACCESSIBILITY REVIEW\n\nThis automated conversion does not certify ADA compliance.\n\nCHANGES\n' + '\n'.join('- '+x for x in changes) + '\n\nHUMAN REVIEW REQUIRED\n' + '\n'.join('- '+x for x in review)
    return output.getvalue(), report
