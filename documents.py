"""Document scanning and semantic text exports. No compliance certification."""
from html import escape
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from converter import convert, MAX_UPLOAD


def text_html(data):
    """Export main-body text, headings, and simple tables in document order.

    Intentionally a text export: visual content and advanced Word constructs
    require review against the original, as stated in the generated HTML.
    """
    doc = Document(BytesIO(data))
    title = doc.core_properties.title or 'Document text'
    language = doc.core_properties.language or 'en-US'
    blocks = []
    for block in doc.iter_inner_content():
        if hasattr(block, 'rows'):
            complex_table = bool(block._tbl.xpath('.//w:gridSpan | .//w:vMerge | .//w:tc/w:tbl'))
            if complex_table:
                blocks.append('<section><h2>Table text — structure requires review</h2>')
                for row in block._tbl.findall(qn('w:tr')):
                    cells = [' '.join(cell.itertext(qn('w:t'))) for cell in row.findall(qn('w:tc'))]
                    blocks.append('<p>' + escape(' | '.join(cells)) + '</p>')
                blocks.append('</section>')
                continue
            blocks.append('<table>')
            for row in block.rows:
                header = bool(row._tr.xpath('./w:trPr/w:tblHeader[not(@w:val="0") and not(@w:val="false")]'))
                blocks.append('<tr>')
                for cell in row.cells:
                    tag = 'th' if header else 'td'
                    attr = ' scope="col"' if header else ''
                    blocks.append(f'<{tag}{attr}>{escape(cell.text).replace(chr(10), "<br>")}</{tag}>')
                blocks.append('</tr>')
            blocks.append('</table>')
            continue
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name
        if style == 'Title':
            # A single page h1 is generated below, avoiding duplicate page titles.
            if text == title:
                continue
            tag = 'h2'
        elif style.startswith('Heading ') and style.split()[-1].isdigit():
            tag = 'h' + str(min(int(style.split()[-1]) + 1, 6))
        else:
            tag = 'p'
        blocks.append(f'<{tag}>{escape(text).replace(chr(10), "<br>")}</{tag}>')
    return (f'<!doctype html><html lang="{escape(language, quote=True)}"><head>'
            '<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{escape(title)}</title><style>'
            'body{font:1.1rem/1.6 Arial,sans-serif;color:#172c3c;background:white;max-width:70ch;margin:2rem auto;padding:0 1rem}'
            'table{border-collapse:collapse;max-width:100%;display:block;overflow-x:auto}th,td{border:1px solid #596976;padding:.5rem;text-align:left}'
            ':focus-visible{outline:3px solid #155d72}aside{border-left:4px solid #155d72;padding-left:1rem}'
            '</style></head><body><main>' + f'<h1>{escape(title)}</h1>'
            '<aside aria-label="Export limitations"><p>This is a text export, not a certified accessible replacement. '
            'Images, headers, footers, footnotes, text boxes, list numbering, and interactive links are not exported. '
            'Compare against the source and review heading levels, table structure, and reading order before sharing.</p></aside>'
            + ''.join(blocks) + '</main></body></html>').encode('utf-8')


def process_document(data, filename, detect_headings=True, table_headers=False, export_pdf=False, high_contrast=True):
    if not data or len(data) > MAX_UPLOAD:
        raise ValueError('Choose a nonempty file smaller than 15 MB.')
    suffix = filename.lower().rsplit('.', 1)[-1]
    if suffix == 'docx':
        word, report = convert(data, detect_headings, table_headers, high_contrast)
        report += '\n\nHTML TEXT EXPORT\nCompare the HTML against the source; images, advanced content, list numbering and link destinations are not exported.\n'
    elif suffix == 'pdf':
        word, report = process_pdf(data, high_contrast)
    else:
        raise ValueError('Choose a .docx or .pdf file. Save older .doc files as .docx in Word first.')
    outputs = {'document': word, 'html': text_html(word)}
    if export_pdf:
        from pdf_export import export_tagged_pdf
        try:
            outputs['pdf'] = export_tagged_pdf(word)
            report += '\n\nTAGGED PDF EXPORT\nCreated a new PDF using LibreOffice PDF/UA export; verified a marked structure tree is present. This is not a full PDF/UA or ADA validation. Review tags, alternative text, reading order, and layout in a PDF accessibility validator.\n'
        except ValueError as exc:
            report += '\n\nPDF EXPORT UNAVAILABLE\n' + str(exc) + '\nWord and HTML downloads remain available.\n'
    from wcag_review import framework_report
    report += framework_report()
    outputs['report'] = report.encode('utf-8')
    return outputs


def process_pdf(data, high_contrast=True):
    if not data.startswith(b'%PDF-'):
        raise ValueError('This file does not appear to be a PDF.')
    try:
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ValueError('Password-protected PDFs are not supported. Upload an unprotected copy.')
        if len(reader.pages) > 200:
            raise ValueError('PDFs are limited to 200 pages. Split this document and upload a smaller section.')
        root = reader.trailer['/Root']
        marked = root.get('/MarkInfo')
        tagged = bool(root.get('/StructTreeRoot')) and bool(marked and marked.get_object().get('/Marked'))
        title = str(reader.metadata.title or 'PDF text export') if reader.metadata else 'PDF text export'
        lang = root.get('/Lang')
        review = ['PDF tags detected; their quality and reading order were not validated.' if tagged else 'The PDF is missing an identifiable marked structure tree. Remediate tags in a PDF accessibility editor.',
                  'Review extracted text against every source page. Tables, columns, lists, images, form fields and mathematical notation may lose meaning during extraction.',
                  'The source PDF is not repaired in place. The app reconstructs a Word text copy and HTML export; optional tagged PDF export uses that reconstructed text and cannot preserve the original layout or visual content.']
        if not lang:
            review.append('The source PDF has no default language metadata.')
        doc = Document(); doc.core_properties.title = title
        if lang:
            doc.core_properties.language = str(lang)
        doc.add_paragraph(title, 'Title')
        total = 0
        for number, page in enumerate(reader.pages, 1):
            # Bound common oversized content streams before extraction.
            content = page.get_contents()
            if content and len(content.get_data()) > 8 * 1024 * 1024:
                raise ValueError('A PDF page is too complex to process. Upload a smaller or simplified PDF.')
            text = page.extract_text() or ''
            total += len(text)
            if total > 2_000_000:
                raise ValueError('The extracted PDF text exceeds the supported limit.')
            doc.add_paragraph(f'Page {number}', 'Heading 1')
            if text.strip():
                for paragraph in text.split('\n\n'):
                    if paragraph.strip():
                        doc.add_paragraph(paragraph.strip())
            else:
                review.append(f'Page {number}: no extractable text. OCR or manual transcription is required; this page is not represented in the text export.')
        if not total:
            raise ValueError('No extractable text was found. This may be a scanned PDF. Run OCR first, then upload the searchable PDF or a Word version.')
        source = BytesIO(); doc.save(source)
        word, base_report = convert(source.getvalue(), False, False, high_contrast)
        report = 'PDF ACCESSIBILITY SCAN\n\nThis automated scan does not certify ADA compliance.\n\nHUMAN REVIEW REQUIRED\n' + '\n'.join('- '+item for item in review) + '\n\nTEXT COPY REMEDIATION\n' + base_report
        return word, report
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('This PDF could not be read. Export a fresh, unprotected PDF and try again.') from exc
